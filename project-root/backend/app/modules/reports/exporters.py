"""Рендер отчётов в файлы: xlsx (openpyxl) и pdf (PyMuPDF + DejaVu Sans).

Используется роутером /reports/export и Celery-задачами app/tasks/reports.py.
Шрифт DejaVuSans.ttf bundled в app/assets/fonts (свободная лицензия) —
без него кириллица в PDF не отображается (встроенные шрифты PyMuPDF
латинские / CJK-only).
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font as XlFont
from openpyxl.utils import get_column_letter

import fitz  # PyMuPDF

from app.modules.reports.schemas import ReportResponse

_FONTS_DIR = Path(__file__).resolve().parent.parent.parent / "assets" / "fonts"

FONT_CANDIDATES = [
    _FONTS_DIR / "DejaVuSans.ttf",
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("C:/Windows/Fonts/arial.ttf"),
]

TEMPLATE_TITLES = {
    "projects": "Отчёт по проектам",
    "tenders": "Отчёт по тендерам",
    "load": "Отчёт по загрузке",
    "finances": "Финансовый отчёт",
    "documents": "Отчёт по документам",
    "remarks": "Отчёт по замечаниям",
}


def _resolve_font_path() -> Path:
    for candidate in FONT_CANDIDATES:
        if candidate.is_file():
            return candidate
    raise RuntimeError(
        "Не найден TTF-шрифт с поддержкой кириллицы. "
        "Ожидался app/assets/fonts/DejaVuSans.ttf"
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def report_to_xlsx_bytes(report: ReportResponse) -> bytes:
    """Собрать .xlsx из отчёта: заголовок, строки, автоширина колонок."""
    wb = Workbook()
    ws = wb.active
    ws.title = report.template[:31]

    columns = list(report.columns)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = XlFont(bold=True)
    for row in report.rows:
        ws.append([_cell_text(row.columns.get(col)) for col in columns])
    ws.freeze_panes = "A2"

    for idx, col in enumerate(columns, start=1):
        max_len = len(str(col))
        for row in report.rows:
            max_len = max(max_len, len(_cell_text(row.columns.get(col))))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 8), 60)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def report_to_pdf_bytes(report: ReportResponse) -> bytes:
    """Собрать .pdf (A4 альбомная, таблица с переносом на новые страницы)."""
    font_path = str(_resolve_font_path())
    font = fitz.Font(fontfile=font_path)

    paper = fitz.paper_rect("a4-l")
    page_w, page_h = paper.width, paper.height
    margin = 36.0
    title_size = 14.0
    header_size = 8.0
    cell_size = 7.5
    header_h = 20.0
    row_h = 14.0

    columns = list(report.columns)
    rows = [[_cell_text(r.columns.get(col)) for col in columns] for r in report.rows]

    # Пропорциональные ширины колонок по длине содержимого (с ограничениями)
    weights = []
    for idx, col in enumerate(columns):
        longest = len(str(col))
        for row in rows:
            longest = max(longest, len(row[idx]))
        weights.append(min(max(longest, 4), 40))
    total_weight = sum(weights)
    avail_w = page_w - margin * 2
    col_widths = [avail_w * w / total_weight for w in weights]

    def truncate(text: str, size: float, width: float) -> str:
        if font.text_length(text, size) <= width - 4:
            return text
        while text and font.text_length(text + "…", size) > width - 4:
            text = text[:-1]
        return text + "…"

    def new_page(doc: fitz.Document):
        page = doc.new_page(width=page_w, height=page_h)
        tw = fitz.TextWriter(page.rect)
        return page, tw

    doc = fitz.open()
    page, tw = new_page(doc)
    y = margin

    title = TEMPLATE_TITLES.get(report.template, f"Отчёт: {report.template}")
    tw.append((margin, y + title_size), title, font=font, fontsize=title_size)
    y += title_size + 6
    generated = report.generated_at.astimezone(timezone.utc).strftime("%d.%m.%Y %H:%M UTC")
    tw.append(
        (margin, y + 9),
        f"Сформировано: {generated} · строк: {len(rows)}",
        font=font,
        fontsize=9,
    )
    y += 24

    def flush_page():
        nonlocal page, tw
        tw.write_text(page)
        page, tw = new_page(doc)

    # Шапка таблицы
    def draw_header():
        nonlocal y
        if y + header_h + row_h > page_h - margin:
            flush_page()
            y = margin
        page.draw_rect(
            fitz.Rect(margin, y, page_w - margin, y + header_h),
            color=None,
            fill=(0.92, 0.92, 0.92),
        )
        x = margin
        for idx, col in enumerate(columns):
            text = truncate(str(col), header_size, col_widths[idx])
            tw.append((x + 3, y + header_h - 6), text, font=font, fontsize=header_size)
            x += col_widths[idx]
        y += header_h

    draw_header()
    for r_idx, row in enumerate(rows):
        if y + row_h > page_h - margin:
            flush_page()
            y = margin
            draw_header()
        if r_idx % 2 == 1:
            page.draw_rect(
                fitz.Rect(margin, y, page_w - margin, y + row_h),
                color=None,
                fill=(0.97, 0.97, 0.97),
            )
        x = margin
        for idx, value in enumerate(row):
            text = truncate(value, cell_size, col_widths[idx])
            tw.append((x + 3, y + row_h - 4), text, font=font, fontsize=cell_size)
            x += col_widths[idx]
        y += row_h

    # Пустой отчёт — честно подписываем
    if not rows:
        tw.append((margin, y + 12), "Нет данных за выбранный период", font=font, fontsize=10)

    tw.write_text(page)
    return doc.tobytes()
